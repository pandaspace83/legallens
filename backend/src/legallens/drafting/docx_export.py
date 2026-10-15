"""Word export (DEC-12, DEC-50).

Clean: what goes to the other side. No source markers, no [UNVERIFIED: …] wrappers, no watermark, no comments,
and empty document properties. The file is re-opened and checked before it is returned; if the check fails, it
is never returned (RISK-03).
Working copy: keeps sources and placeholders, and every page is stamped as privileged work product.
"""

import io
import re
import zipfile
from datetime import date
from typing import Any

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor

from legallens.drafting.placeholders import PLACEHOLDER_RE, strip_citation_markers, unwrap_unverified

WORK_PRODUCT_STAMP = "PRIVILEGED & CONFIDENTIAL / ATTORNEY WORK PRODUCT / DO NOT SEND"
WATERMARK = "AI DRAFT · NOT REVIEWED"
FRE408_HEADER = "CONFIDENTIAL SETTLEMENT COMMUNICATION: FRE 408 AND STATE EQUIVALENTS"

_EMPTY_CORE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
    'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"/>'
)
_EMPTY_APP = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"/>'
)
_FORBIDDEN_IN_CLEAN = re.compile(
    r"\[(PLACEHOLDER|CITE NEEDED|RULE NEEDED|UNVERIFIED)|\[\d{1,3}\]|AI DRAFT|WORK PRODUCT|PRIVILEGED"
)


class ExportVerificationError(RuntimeError):
    pass


def reservation_text(client_name: str) -> str:
    return (
        f"Nothing in this letter is intended as, or should be read as, a waiver of any of {client_name}'s rights, "
        "claims or defenses, all of which are expressly reserved."
    )


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


def _add_text(doc: Any, text: str, *, bold: bool = False) -> None:
    for block in _paragraphs(text):
        para = doc.add_paragraph()
        lines = block.split("\n")
        for i, line in enumerate(lines):
            run = para.add_run(line)
            run.bold = bold
            if i < len(lines) - 1:
                run.add_break()


def build(draft: dict[str, Any], *, kind: str, client_name: str, today: date | None = None) -> bytes:
    """draft: {type, status, sections: [{id, heading, content, citations, kind}], safeguards, citations_index}."""
    clean = kind == "clean"
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    def text_of(value: str) -> str:
        return unwrap_unverified(strip_citation_markers(value)) if clean else value

    if not clean:
        header = doc.sections[0].header.paragraphs[0]
        header.text = WORK_PRODUCT_STAMP
        if draft["status"] != "reviewed_adopted":
            run = doc.add_paragraph().add_run(WATERMARK)
            run.bold = True
            run.font.color.rgb = RGBColor(0xA8, 0x26, 0x1F)

    safeguards = draft.get("safeguards") or {}
    if safeguards.get("fre408_header"):
        _add_text(doc, FRE408_HEADER, bold=True)
    day = today or date.today()
    doc.add_paragraph(draft.get("date_line") or f"{day:%B} {day.day}, {day.year}")

    closing_index = next((i for i, s in enumerate(draft["sections"]) if s.get("kind") == "closing"), None)
    for index, section in enumerate(draft["sections"]):
        if index == closing_index and safeguards.get("reservation_of_rights"):
            _add_text(doc, reservation_text(client_name))
        if section.get("show_heading", True) and section.get("heading"):
            _add_text(doc, text_of(section["heading"]), bold=True)
        _add_text(doc, text_of(section.get("content", "")))
    if closing_index is None and safeguards.get("reservation_of_rights"):
        _add_text(doc, reservation_text(client_name))

    if not clean:
        sources = [c for s in draft["sections"] for c in s.get("citations", [])]
        if sources:
            _add_text(doc, "Sources", bold=True)
            for c in sorted({c["n"]: c for c in sources}.values(), key=lambda c: c["n"]):
                where = f"p. {c['page']}" + (f", {c['para_ref']}" if c.get("para_ref") else "")
                doc.add_paragraph(f"[{c['n']}] {c.get('document_name') or c['document_id']}, {where}")

    buffer = io.BytesIO()
    doc.save(buffer)
    data = _scrub_package(buffer.getvalue()) if clean else _scrub_author(buffer.getvalue())
    if clean:
        verify_clean(data)
    return data


def build_summary(summary: dict[str, Any]) -> bytes:
    """A summary as a privileged working copy (decision 5, DEC-12): it keeps its source numbers and lists them, so
    every statement can be checked, and it is stamped as work product. Summaries are never sent, so there is no
    clean copy.

    summary: {document_name, template_title, audience_title, coverage_text, sections: [{title, found, content,
    citations}], model_id}."""
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    doc.sections[0].header.paragraphs[0].text = WORK_PRODUCT_STAMP
    _add_text(doc, f"Summary: {summary['document_name']}", bold=True)
    doc.add_paragraph(f"{summary['template_title']} · {summary['audience_title']}")
    if summary.get("coverage_text"):
        doc.add_paragraph(f"Coverage: {summary['coverage_text']}")
    for section in summary["sections"]:
        _add_text(doc, section["title"], bold=True)
        _add_text(doc, section["content"])
    sources = [c for s in summary["sections"] for c in s.get("citations", [])]
    if sources:
        _add_text(doc, "Sources", bold=True)
        for c in sorted({c["n"]: c for c in sources}.values(), key=lambda c: c["n"]):
            where = f"p. {c['page']}" + (f", {c['para_ref']}" if c.get("para_ref") else "")
            doc.add_paragraph(f"[{c['n']}] {c.get('document_name') or c['document_id']}, {where}")
    doc.add_paragraph(f"AI summary, checked against the document's text. {summary.get('model_id') or ''}".strip())
    buffer = io.BytesIO()
    doc.save(buffer)
    return _scrub_author(buffer.getvalue())


def _rewrite_zip(data: bytes, replacements: dict[str, str], drop: set[str]) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            if item.filename in drop:
                continue
            payload = replacements[item.filename].encode() if item.filename in replacements else src.read(item.filename)
            dst.writestr(item, payload)
    return out.getvalue()


def _scrub_package(data: bytes) -> bytes:
    """Empty core and app properties; no comments or custom properties parts."""
    return _rewrite_zip(
        data,
        {"docProps/core.xml": _EMPTY_CORE, "docProps/app.xml": _EMPTY_APP},
        {"word/comments.xml", "docProps/custom.xml"},
    )


def _scrub_author(data: bytes) -> bytes:
    return _rewrite_zip(data, {"docProps/core.xml": _EMPTY_CORE}, set())


def verify_clean(data: bytes) -> None:
    package = zipfile.ZipFile(io.BytesIO(data))
    names = set(package.namelist())
    if names & {"word/comments.xml", "docProps/custom.xml"}:
        raise ExportVerificationError("comments or custom properties present")
    doc = DocxDocument(io.BytesIO(data))
    props = doc.core_properties
    for field in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
        if getattr(props, field):
            raise ExportVerificationError(f"document property {field} is set")
    texts = [p.text for p in doc.paragraphs]
    for section in doc.sections:
        texts.extend(p.text for p in section.header.paragraphs)
        texts.extend(p.text for p in section.footer.paragraphs)
    full = "\n".join(texts)
    if _FORBIDDEN_IN_CLEAN.search(full) or PLACEHOLDER_RE.search(full):
        raise ExportVerificationError("internal markers remain in the text")
