"""F · Summaries, G · Drafts (versions, checks, review, export) and the matter export."""

import io
import zipfile

import pytest
from docx import Document as DocxDocument

from tests.conftest import Api


@pytest.fixture
def smith(api: Api) -> dict:  # type: ignore[type-arg]
    matter = api.create_matter(
        profile={
            "parties": [
                {
                    "name": "Acme Corp",
                    "role": "opposing",
                    "represented": True,
                    "counsel": {
                        "name": "Dana Whitfield",
                        "firm": "Whitfield & Grant LLP",
                        "address": "100 Market Street, San Francisco, CA 94105",
                    },
                }
            ]
        },
    )
    ids = api.ingest(matter["id"], "smith_msa_v3.pdf", "acme_demand_letter_scanned.pdf", "exhibit_b_invoices.pdf")
    return {**matter, "msa": ids[0], "letter": ids[1], "exhibit": ids[2]}


def _draft(api: Api, smith: dict) -> dict:  # type: ignore[type-arg]
    chat = api.new_chat(smith["id"])
    turn = api.ask(
        chat,
        "Draft a response letter to the demand",
        mode="draft",
        options={"incoming_document_id": smith["letter"], "supporting_document_ids": [smith["exhibit"]]},
    )
    draft_id = turn.first("job")["entity_id"]
    assert api.client.get(f"/api/v1/drafts/{draft_id}").json()["status"] == "generating"
    api.run_jobs()
    return api.client.get(f"/api/v1/drafts/{draft_id}").json()  # type: ignore[no-any-return]


class TestLists:
    """INT-1.3: a matter's summaries and drafts, newest first, and richer matter list items."""

    def test_summaries_newest_first_with_document_and_chat(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        first = api.client.post(
            f"/api/v1/documents/{smith['letter']}/summaries", json={"template": "letter_demand"}
        ).json()["summary_id"]
        chat = api.new_chat(smith["id"])
        second = api.ask(chat, "Summarize the exhibit_b_invoices in one paragraph").first("job")["entity_id"]
        url = f"/api/v1/matters/{smith['id']}/summaries"
        items = api.client.get(url).json()["items"]
        assert [s["id"] for s in items] == [second, first]
        assert (items[0]["conversation_id"], items[1]["conversation_id"]) == (chat, None)
        assert items[1]["document_name"] == "acme_demand_letter_scanned.pdf" and items[1]["status"] == "pending"
        only = api.client.get(url, params={"document_id": smith["letter"]}).json()["items"]
        assert [s["id"] for s in only] == [first]
        page = api.client.get(url, params={"limit": 1}).json()
        rest = api.client.get(url, params={"limit": 1, "cursor": page["next_cursor"]}).json()
        assert [page["items"][0]["id"], rest["items"][0]["id"]] == [second, first] and rest["next_cursor"] is None

    def test_drafts_list(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        draft = _draft(api, smith)
        (item,) = api.client.get(f"/api/v1/matters/{smith['id']}/drafts").json()["items"]
        assert item["id"] == draft["id"] and item["title"] == draft["title"]
        assert item["incoming_document_name"] == "acme_demand_letter_scanned.pdf"
        assert (item["status"], item["version"], item["conversation_id"]) == (
            "not_reviewed",
            draft["version"],
            draft["conversation_id"],
        )

    def test_matter_list_items_name_the_other_side(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        api.client.patch(f"/api/v1/matters/{smith['id']}", json={"court": "N.D. Cal.", "practice_area": "Contracts"})
        bare = api.create_matter("Bare matter")
        items = {m["id"]: m for m in api.client.get("/api/v1/matters").json()["items"]}
        mine = items[smith["id"]]
        assert (mine["court"], mine["practice_area"], mine["opposing_party"]) == ("N.D. Cal.", "Contracts", "Acme Corp")
        assert mine["opposing_counsel"] == {"name": "Dana Whitfield", "firm": "Whitfield & Grant LLP"}
        assert items[bare["id"]]["opposing_party"] is None and items[bare["id"]]["opposing_counsel"] is None


class TestSummaryExport:
    """INT-1.7: the summary names its document and chat, and exports as a .docx working copy (decision 5)."""

    def test_working_copy_with_sources(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        chat = api.new_chat(smith["id"])
        summary_id = api.ask(chat, "Summarize the acme_demand_letter_scanned in detail").first("job")["entity_id"]
        early = api.client.post(f"/api/v1/summaries/{summary_id}/export")
        assert (early.status_code, early.json()["code"]) == (409, "summary_not_ready")
        api.run_jobs()
        summary = api.client.get(f"/api/v1/summaries/{summary_id}").json()
        assert summary["document_name"] == "acme_demand_letter_scanned.pdf" and summary["conversation_id"] == chat

        exported = api.client.post(f"/api/v1/summaries/{summary_id}/export")
        assert exported.status_code == 200, exported.text
        assert (
            exported.headers["content-disposition"] == 'attachment; filename="summary-acme-demand-letter-scanned.docx"'
        )
        assert exported.headers["cache-control"] == "private, no-store"
        doc = DocxDocument(io.BytesIO(exported.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert doc.sections[0].header.paragraphs[0].text.startswith("PRIVILEGED & CONFIDENTIAL")
        assert "Summary: acme_demand_letter_scanned.pdf" in text and "Detailed digest" in text
        assert "Sources" in text and "[1]" in text  # every statement stays checkable
        assert not doc.core_properties.author  # no author left in the file


class TestSummaries:
    def test_letter_template_with_coverage(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        accepted = api.client.post(
            f"/api/v1/documents/{smith['letter']}/summaries", json={"template": "letter_demand", "audience": "detailed"}
        ).json()
        assert api.client.get(f"/api/v1/summaries/{accepted['summary_id']}").json()["status"] == "pending"
        api.run_jobs()
        summary = api.client.get(f"/api/v1/summaries/{accepted['summary_id']}").json()
        assert summary["status"] == "ready" and summary["model_id"] == "stub-orchestrator"
        fields = {s["field"]: s for s in summary["sections"]}
        assert list(fields) == ["sender", "recipient", "demands", "deadlines", "claims", "threats", "amounts"]
        assert fields["demands"]["found"] and fields["demands"]["citations"][0]["document_id"] == smith["letter"]
        assert "$48,500" in fields["amounts"]["content"] or "$48,500" in fields["demands"]["content"]
        assert summary["coverage"]["text"] == "1 of 1 page read; OCR on p. 1"
        assert api.client.get(f"/api/v1/jobs/{accepted['job_id']}").json()["result"]["kind"] == "summary"

    def test_missing_fields_say_not_found(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        accepted = api.client.post(
            f"/api/v1/documents/{smith['exhibit']}/summaries", json={"template": "opinion", "audience": "one_paragraph"}
        ).json()
        api.run_jobs()
        sections = api.client.get(f"/api/v1/summaries/{accepted['summary_id']}").json()["sections"]
        missing = [s for s in sections if not s["found"]]
        assert missing and all(s["content"] == "Not found in this document." for s in missing)

    def test_not_ready_document(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        pending = api.upload(smith["id"], "rivera_lease.txt").json()
        response = api.client.post(
            f"/api/v1/documents/{pending['document_id']}/summaries", json={"template": "letter_demand"}
        )
        assert (response.status_code, response.json()["code"]) == (409, "document_not_ready")


class TestDrafts:
    def test_first_version_structure_and_safeguards(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        draft = _draft(api, smith)
        assert draft["status"] == "not_reviewed" and draft["watermark"] == "AI draft · not reviewed"
        assert draft["version"] == 1 and [v["source"] for v in draft["versions"]] == ["ai"]
        assert draft["safeguards"]["addressee"] == {
            "name": "Dana Whitfield",
            "firm": "Whitfield & Grant LLP",
            "address": "100 Market Street, San Francisco, CA 94105",
            "email": None,
            "rule": "4.2",
        }
        assert not draft["safeguards"]["fre408_header"] and not draft["safeguards"]["reservation_of_rights"]
        points = [s for s in draft["sections"] if s["claim_ids"]]
        assert len(points) >= 3 and all("[PLACEHOLDER:" in s["content"] for s in points)
        kinds = {p["kind"] for p in draft["placeholders"]}
        assert {"placeholder", "cite_needed"} <= kinds

    def test_checks_flag_unverified_authority_from_the_other_side(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        draft = _draft(api, smith)
        checks = api.client.get(f"/api/v1/drafts/{draft['id']}/checks").json()
        assert {p["status"] for p in checks["point_by_point"]} == {"flagged"}
        case = next(i for i in checks["citation_checklist"] if i["kind"] == "case")
        assert case["authority"].startswith("Acme Corp. v. Beta Transport")
        assert (case["status"], case["unverified"], case["acknowledged"]) == ("only_mentioned", True, False)
        assert case["supporting"]["document_id"] == smith["letter"]
        assert "citator" in case["note"]
        body = " ".join(s["content"] for s in draft["sections"])
        assert "[UNVERIFIED: Acme Corp. v. Beta Transport" in body  # DEC-48
        codes = {b["code"] for b in checks["export"]["blockers"]}
        assert codes == {"placeholders", "unacknowledged_citations", "not_reviewed"}

    def test_edit_review_acknowledge_and_export_clean(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        draft = _draft(api, smith)
        url = f"/api/v1/drafts/{draft['id']}"
        blocked = api.client.post(f"{url}/export", params={"kind": "clean"})
        assert blocked.status_code == 409 and blocked.json()["code"] == "export_blocked"

        stale = api.client.patch(url, json={"base_version": 0, "sections": []})
        assert (stale.status_code, stale.json()["code"]) == (409, "version_conflict")

        edits = []
        for section in draft["sections"]:
            content = section["content"]
            for p in [p for p in draft["placeholders"] if p["section_id"] == section["id"]]:
                if p["kind"] != "unverified":
                    content = content.replace(p["token"], "Smith disputes this point.")
            edits.append({"id": section["id"], "content": content})
        edited = api.client.patch(
            url,
            json={
                "base_version": 1,
                "sections": edits,
                "safeguards": {"fre408_header": True, "reservation_of_rights": True},
            },
        )
        assert edited.status_code == 200, edited.text
        edited = edited.json()
        assert edited["version"] == 2 and edited["sections"][0]["origin"] == "lawyer"
        assert {p["kind"] for p in edited["placeholders"]} == {"unverified"}

        checks = api.client.get(f"{url}/checks").json()
        assert {p["status"] for p in checks["point_by_point"]} == {"addressed"}
        ack = [i["id"] for i in checks["citation_checklist"]]
        api.client.patch(url, json={"base_version": 2, "acknowledge": ack})
        assert api.client.get(f"{url}/checks").json()["export"]["blockers"] == [
            {"code": "not_reviewed", "message": 'Mark the draft "Reviewed and adopted" first.'}
        ]
        reviewed = api.client.post(f"{url}/review", json={"confirm": True}).json()
        assert reviewed["status"] == "reviewed_adopted" and reviewed["reviewed_at"] and reviewed["watermark"] is None

        clean = api.client.post(f"{url}/export", params={"kind": "clean"})
        assert clean.status_code == 200, clean.text
        assert clean.headers["content-disposition"] == 'attachment; filename="response-letter-v2.docx"'
        doc = DocxDocument(io.BytesIO(clean.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "CONFIDENTIAL SETTLEMENT COMMUNICATION" in text and "expressly reserved" in text
        assert "Acme Corp. v. Beta Transport" in text  # acknowledged, so written as cited
        for marker in ("[UNVERIFIED", "[PLACEHOLDER", "[1]", "AI DRAFT", "PRIVILEGED"):
            assert marker not in text
        props = doc.core_properties
        assert not any([props.author, props.last_modified_by, props.title, props.comments])
        assert "word/comments.xml" not in zipfile.ZipFile(io.BytesIO(clean.content)).namelist()

        working = api.client.post(f"{url}/export", params={"kind": "working"})
        working_doc = DocxDocument(io.BytesIO(working.content))
        assert "ATTORNEY WORK PRODUCT" in working_doc.sections[0].header.paragraphs[0].text
        assert "Sources" in "\n".join(p.text for p in working_doc.paragraphs)

        # Editing an adopted draft needs a fresh review (DEC-51).
        again = api.client.patch(
            url, json={"base_version": 2, "sections": [{"id": "s3", "content": "Dear Ms. Whitfield:"}]}
        )
        assert again.json()["status"] == "not_reviewed" and again.json()["watermark"]

    def test_regenerate_restore_and_refine_from_chat(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        draft = _draft(api, smith)
        url = f"/api/v1/drafts/{draft['id']}"
        section = next(s for s in draft["sections"] if s["claim_ids"])
        job = api.client.post(f"{url}/sections/{section['id']}/regenerate", json={"instruction": "Make it firmer"})
        assert job.status_code == 202
        api.run_jobs()
        regenerated = api.client.get(url).json()
        assert regenerated["version"] == 2 and regenerated["versions"][-1]["source"] == "regenerate"
        assert regenerated["last_diff"]["section_id"] == section["id"]
        assert any(op["op"] == "insert" and "firmer" in op["text"] for op in regenerated["last_diff"]["ops"])

        old = api.client.get(url, params={"version": 1}).json()
        assert old["version"] == 1 and not old["is_current"]
        restored = api.client.patch(url, json={"base_version": 2, "restore_version": 1}).json()
        assert restored["version"] == 3 and restored["sections"] == old["sections"]

        chat = draft["conversation_id"]
        turn = api.ask(chat, "Make point 2 firmer", options={"draft_id": draft["id"]})
        assert turn.first("intent")["intent"] == "refine_draft" and "version 4" in turn.answer["content"]
        assert api.client.get(url).json()["version"] == 4

    def test_rule_4_3_reminder_for_unrepresented_party(self, api: Api) -> None:
        matter = api.create_matter(
            "Rivera", profile={"parties": [{"name": "Northgate Properties", "role": "opposing"}]}
        )
        [incoming] = api.ingest(matter["id"], "acme_demand_letter_scanned.pdf")
        chat = api.new_chat(matter["id"])
        turn = api.ask(chat, "Write a reply to this letter", options={"incoming_document_id": incoming})
        api.run_jobs()
        draft = api.client.get(f"/api/v1/drafts/{turn.first('job')['entity_id']}").json()
        assert draft["safeguards"]["addressee"]["rule"] == "4.3" and draft["safeguards"]["rule_4_3_reminder"]

    def test_memo_has_rule_needed_blocks(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        chat = api.new_chat(smith["id"])
        turn = api.ask(
            chat,
            "Draft a memo on the late delivery claim",
            mode="draft",
            options={"incoming_document_id": smith["letter"], "draft_type": "memo"},
        )
        api.run_jobs()
        draft = api.client.get(f"/api/v1/drafts/{turn.first('job')['entity_id']}").json()
        assert [s["heading"] for s in draft["sections"]] == [
            "Memorandum",
            "Question Presented",
            "Brief Answer",
            "Statement of Facts",
            "Discussion",
            "Conclusion",
        ]
        assert any(p["kind"] == "rule_needed" for p in draft["placeholders"])


class TestDraftIntegration:
    """INT-1.6: what the frontend's draft screens need from the server."""

    def test_start_a_draft_directly(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        url = f"/api/v1/matters/{smith['id']}/drafts"
        accepted = api.client.post(
            url,
            json={
                "type": "response_letter",
                "incoming_document_id": smith["letter"],
                "supporting_document_ids": [smith["msa"], smith["exhibit"]],
                "instructions": "Dispute the invoices.",
            },
        )
        assert accepted.status_code == 202, accepted.text
        api.run_jobs()
        draft = api.client.get(f"/api/v1/drafts/{accepted.json()['draft_id']}").json()
        assert (draft["status"], draft["conversation_id"], draft["model_id"]) == (
            "not_reviewed",
            None,
            "stub-orchestrator",
        )
        assert draft["date_line"] and draft["versions"][0]["label"] == "AI draft"
        assert set(draft["supporting_document_ids"]) == {smith["msa"], smith["exhibit"]}
        head = {s["kind"]: s for s in draft["sections"][:3]}
        assert set(head) == {"addressee", "re", "intro"} and not head["addressee"]["show_heading"]
        other = api.create_matter("Other")
        bad = api.client.post(
            f"/api/v1/matters/{other['id']}/drafts", json={"type": "memo", "incoming_document_id": smith["letter"]}
        )
        assert (bad.status_code, bad.json()["code"]) == (422, "invalid_option")  # another matter's document

    def test_memo_header_and_version_labels(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        accepted = api.client.post(
            f"/api/v1/matters/{smith['id']}/drafts", json={"type": "memo", "incoming_document_id": smith["letter"]}
        ).json()
        api.run_jobs()
        url = f"/api/v1/drafts/{accepted['draft_id']}"
        memo = api.client.get(url).json()
        assert memo["memo_header"] == {"to": "[RECIPIENT]", "from": "Dev Lawyer", "re": "Smith v. Acme Corp"}
        head = memo["sections"][0]
        edited = api.client.patch(
            url,
            json={
                "base_version": 1,
                "sections": [{"id": head["id"], "content": head["content"].replace("[RECIPIENT]", "File")}],
                "version_label": "Addressed it to the file",
            },
        ).json()
        assert edited["memo_header"]["to"] == "File" and edited["versions"][-1]["label"] == "Addressed it to the file"
        restored = api.client.patch(url, json={"base_version": 2, "restore_version": 1}).json()
        assert (
            restored["versions"][-1]["label"] == "Restored version 1" and restored["memo_header"]["to"] == "[RECIPIENT]"
        )

    def test_rewrite_is_a_preview_until_accepted(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        draft = _draft(api, smith)
        url = f"/api/v1/drafts/{draft['id']}"
        point = next(s for s in draft["sections"] if s["kind"] == "point")
        preview = api.client.post(f"{url}/sections/{point['id']}/propose", json={"instruction": "Make it firmer"})
        assert preview.status_code == 200, preview.text
        preview = preview.json()
        assert preview["body"] != point["content"] and any(op["op"] == "insert" for op in preview["diff"])
        unchanged = api.client.get(url).json()
        assert unchanged["version"] == 1 and unchanged["sections"] == draft["sections"]  # nothing saved yet

        accepted = api.client.patch(url, json={"base_version": 1, "accept_proposal": preview["proposal_id"]}).json()
        section = next(s for s in accepted["sections"] if s["id"] == point["id"])
        assert accepted["version"] == 2 and section["content"] == preview["body"] and section["origin"] == "ai"
        assert accepted["versions"][-1]["source"] == "regenerate"
        assert accepted["versions"][-1]["label"].endswith("Make it firmer")
        again = api.client.patch(url, json={"base_version": 2, "accept_proposal": preview["proposal_id"]})
        assert again.status_code == 404  # used up

        stale = api.client.post(f"{url}/sections/{point['id']}/propose", json={"instruction": "Shorter"}).json()
        api.client.patch(url, json={"base_version": 2, "sections": [{"id": point["id"], "content": "Edited."}]})
        late = api.client.patch(url, json={"base_version": 3, "accept_proposal": stale["proposal_id"]})
        assert (late.status_code, late.json()["code"]) == (409, "proposal_stale")

    def test_claims_waive_and_answer(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        draft = _draft(api, smith)
        url = f"/api/v1/drafts/{draft['id']}"
        claims = draft["claims"]
        assert claims and all(c["status"] == "flagged" and c["addressed_by"] for c in claims)
        assert claims[0]["source"]["document_id"] == smith["letter"]
        waived = api.client.patch(url, json={"base_version": 1, "waived_claims": [claims[0]["id"]]}).json()
        assert waived["version"] == 1 and waived["claims"][0]["waived"]  # no new version
        bad = api.client.patch(url, json={"base_version": 1, "waived_claims": ["c99"]})
        assert (bad.status_code, bad.json()["code"]) == (422, "unknown_claim")

        preview = api.client.post(f"{url}/claims/{claims[1]['id']}/propose").json()
        new = preview["section"]
        assert new["claim_ids"] == [claims[1]["id"]] and "[PLACEHOLDER:" in new["content"]
        numbers = [c["n"] for s in draft["sections"] for c in s["citations"]]
        assert new["citations"][0]["n"] == max(numbers) + 1  # its source number doesn't clash
        added = api.client.patch(url, json={"base_version": 1, "accept_proposal": preview["proposal_id"]}).json()
        ids = [s["id"] for s in added["sections"]]
        assert ids.index(new["id"]) == len(ids) - 2  # before the closing
        assert added["versions"][-1]["label"].startswith("Answered:")

    def test_named_sources_and_acknowledgements(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        draft = _draft(api, smith)
        url = f"/api/v1/drafts/{draft['id']}"
        named = api.client.patch(
            url, json={"base_version": 1, "added_authorities": ["42 U.S.C. § 1983", "SLA incident log, Aug 2026"]}
        ).json()
        assert named["added_authorities"] == ["42 U.S.C. § 1983", "SLA incident log, Aug 2026"]
        rows = {i["authority"]: i for i in api.client.get(f"{url}/checks").json()["citation_checklist"]}
        assert rows["42 U.S.C. § 1983"]["status"] == "added_by_lawyer"
        assert "SLA incident log, Aug 2026" not in rows  # evidence, not authority (decision 1)
        item = rows["42 U.S.C. § 1983"]["id"]
        assert item in api.client.patch(url, json={"base_version": 1, "acknowledge": [item]}).json()["acknowledged"]
        assert (
            item not in api.client.patch(url, json={"base_version": 1, "unacknowledge": [item]}).json()["acknowledged"]
        )

    def test_refusals(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        url = f"/api/v1/matters/{smith['id']}/drafts"
        body = {"type": "memo", "incoming_document_id": smith["letter"]}
        pending = api.client.post(url, json=body).json()["draft_id"]
        early = api.client.post(f"/api/v1/drafts/{pending}/claims/c1/propose")
        assert (early.status_code, early.json()["code"]) == (409, "draft_not_ready")
        for extra, code in (
            ({"supporting_document_ids": ["doc_nope"]}, "invalid_option"),
            ({"conversation_id": "cnv_nope"}, "invalid_option"),
        ):
            bad = api.client.post(url, json={**body, **extra})
            assert (bad.status_code, bad.json()["code"]) == (422, code)
        api.run_jobs()

        draft = _draft(api, smith)
        durl = f"/api/v1/drafts/{draft['id']}"
        assert api.client.post(f"{durl}/sections/s99/propose", json={"instruction": "x"}).status_code == 404
        assert api.client.post(f"{durl}/claims/c99/propose").status_code == 404
        both = api.client.patch(durl, json={"base_version": 1, "restore_version": 1, "accept_proposal": "prp_x"})
        assert (both.status_code, both.json()["code"]) == (422, "invalid_option")
        ids = [
            api.client.post(f"{durl}/sections/s3/propose", json={"instruction": f"try {i}"}).json()["proposal_id"]
            for i in range(21)
        ]
        expired = api.client.patch(durl, json={"base_version": 1, "accept_proposal": ids[0]})
        assert expired.status_code == 404  # only the newest 20 previews are kept
        assert api.client.patch(durl, json={"base_version": 1, "accept_proposal": ids[-1]}).status_code == 200

        page = api.client.get(url, params={"limit": 1}).json()
        rest = api.client.get(url, params={"limit": 1, "cursor": page["next_cursor"]}).json()
        assert page["items"][0]["id"] != rest["items"][0]["id"]

    def test_review_needs_everything_resolved_and_changes_undo_it(self, api: Api, smith: dict) -> None:  # type: ignore[type-arg]
        draft = _draft(api, smith)
        url = f"/api/v1/drafts/{draft['id']}"
        blocked = api.client.post(f"{url}/review", json={"confirm": True})
        assert (blocked.status_code, blocked.json()["code"]) == (409, "review_blocked")
        assert {b["code"] for b in blocked.json()["blockers"]} == {"placeholders", "unacknowledged_citations"}
        resolved = api.resolve_draft(draft["id"])
        adopted = api.client.post(f"{url}/review", json={"confirm": True}).json()
        assert adopted["status"] == "reviewed_adopted"
        switched = api.client.patch(
            url, json={"base_version": resolved["version"], "safeguards": {"fre408_header": True}}
        ).json()
        assert switched["status"] == "not_reviewed" and switched["watermark"]  # the letter changed (DEC-51)


def test_matter_export_zip(api: Api, smith: dict) -> None:  # type: ignore[type-arg]
    _draft(api, smith)
    accepted = api.client.post(f"/api/v1/matters/{smith['id']}/export")
    assert accepted.status_code == 202
    job_id = accepted.json()["job_id"]
    assert api.client.get(f"/api/v1/jobs/{job_id}/file").json()["code"] == "job_not_done"
    api.run_jobs()
    response = api.client.get(f"/api/v1/jobs/{job_id}/file")
    assert response.status_code == 200 and response.headers["content-type"] == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(response.content)).namelist()
    assert {"matter.json", "manifest.json", "audit.jsonl"} <= set(names)
    assert sum(n.startswith("documents/") and n.endswith(".pdf") for n in names) == 3
    assert any(n.endswith("-working-copy.docx") for n in names)
    assert any(n.startswith("conversations/") for n in names)


def test_cancel_a_queued_job(api: Api, smith: dict) -> None:  # type: ignore[type-arg]
    job_id = api.client.post(f"/api/v1/matters/{smith['id']}/export").json()["job_id"]
    cancelled = api.client.post(f"/api/v1/jobs/{job_id}/cancel").json()
    assert cancelled["status"] == "cancelled" and cancelled["error"]["code"] == "cancelled"
    again = api.client.post(f"/api/v1/jobs/{job_id}/cancel")
    assert (again.status_code, again.json()["code"]) == (409, "job_not_cancellable")
    assert api.run_jobs() == 0


OPINION = """ACME CORP. v. BETA TRANSPORT

United States Court of Appeals for the Ninth Circuit. Argued and submitted May 1, 2008. Filed June 2, 2008.

Acme Corp. v. Beta Transport, 512 F.3d 1201 (9th Cir. 2008).

OPINION

Late performance of a freight contract is a material breach when time is of the essence. Affirmed.
"""


def test_uploaded_opinion_counts_as_full_text(api: Api, smith: dict) -> None:  # type: ignore[type-arg]
    response = api.upload(smith["id"], "acme_v_beta_opinion.txt", "client", OPINION.encode())
    assert response.status_code == 202
    api.run_jobs()
    draft = _draft(api, smith)
    checks = api.client.get(f"/api/v1/drafts/{draft['id']}/checks").json()
    case = next(i for i in checks["citation_checklist"] if i["kind"] == "case")
    assert (case["status"], case["unverified"]) == ("full_text_uploaded", False)
    assert case["supporting"]["document_id"] == response.json()["document_id"]
    body = " ".join(s["content"] for s in draft["sections"])
    assert "[UNVERIFIED: Acme Corp. v. Beta Transport" not in body
